# build_data.py
# Reads DRBA_Calendar.xlsx and writes the dataset the web view runs on:
#   data.js    window.DRBA_DATA - loaded by index.html, works over file://
#   data.json  the same records, for anything else that wants them
#   data.csv   UTF-8 with BOM so Excel on Windows keeps the diacritics
# Source of truth is the lunar_ref sheet (Hong Kong Observatory transcription)
# plus the dharma events typed off the DRBA 2026 printed calendar.
# Version 1.0

import csv
import json
import datetime
import openpyxl

XLSX = "DRBA_Calendar.xlsx"
OUT = "."

DAY_CN = ["初一", "初二", "初三", "初四", "初五", "初六", "初七", "初八",
          "初九", "初十", "十一", "十二", "十三", "十四", "十五", "十六",
          "十七", "十八", "十九", "二十", "廿一", "廿二", "廿三", "廿四",
          "廿五", "廿六", "廿七", "廿八", "廿九", "三十"]
DAY_PY = ["chu yi", "chu er", "chu san", "chu si", "chu wu", "chu liu",
          "chu qi", "chu ba", "chu jiu", "chu shi", "shi yi", "shi er",
          "shi san", "shi si", "shi wu", "shi liu", "shi qi", "shi ba",
          "shi jiu", "er shi", "nian yi", "nian er", "nian san", "nian si",
          "nian wu", "nian liu", "nian qi", "nian ba", "nian jiu",
          "san shi"]
MONTH_CN = ["正月", "二月", "三月", "四月", "五月", "六月", "七月", "八月",
            "九月", "十月", "十一月", "十二月"]
MONTH_PY = ["zheng yue", "er yue", "san yue", "si yue", "wu yue", "liu yue",
            "qi yue", "ba yue", "jiu yue", "shi yue", "shi yi yue",
            "shi er yue"]
ORD = {1: "1st", 2: "2nd", 3: "3rd", 4: "4th", 5: "5th", 6: "6th", 7: "7th",
       8: "8th", 9: "9th", 10: "10th", 11: "11th", 12: "12th"}

DAY_NUM = {n: i + 1 for i, n in enumerate(DAY_CN)}
DAY_NUM["三 十"] = 30                     # the DRBA calendar's spaced variant
MONTH_NUM = {n: i + 1 for i, n in enumerate(MONTH_CN)}

print("[*] reading %s" % XLSX)
wb = openpyxl.load_workbook(XLSX, data_only=True)

lr = wb["lunar_ref"]
rows = []
r = 2
while lr.cell(r, 1).value is not None:
    d = lr.cell(r, 1).value
    rows.append((d.date() if hasattr(d, "date") else d,
                 lr.cell(r, 2).value, lr.cell(r, 3).value))
    r += 1
print("    %d reference days, %s to %s"
      % (len(rows), rows[0][0], rows[-1][0]))

# Dharma events, typed off the DRBA printed calendar. The date column on that
# tab is a formula chain with no cached values, so dates come from the row
# offset instead: row 2 is the first day of the reference range.
events = {}
drba = openpyxl.load_workbook(XLSX)["drba"]
first = rows[0][0]
for r in range(2, 175):
    ev = drba.cell(r, 7).value
    if ev:
        events[first + datetime.timedelta(days=r - 2)] = str(ev).strip()
drba_last = first + datetime.timedelta(days=172)
print("    %d dharma events, DRBA range ends %s" % (len(events), drba_last))

# month lengths: distance from each day-1 to the next
starts = [i for i, (_, _, day) in enumerate(rows) if DAY_NUM[day] == 1]
length_at = {}
for k, i in enumerate(starts):
    end = starts[k + 1] if k + 1 < len(starts) else len(rows)
    for j in range(i, end):
        length_at[j] = end - i
print("    %d lunar months spanned" % len(starts))

print("[*] applying the rules")
records = []
for i, (d, month_cn, day_cn) in enumerate(rows):
    day = DAY_NUM[day_cn]
    leap = month_cn.startswith("閏")
    base = month_cn[1:] if leap else month_cn
    mnum = MONTH_NUM[base]
    mlen = length_at.get(i)
    # Six fast days: 8, 14, 15, 23 and the last two days of the month. The
    # last two are found by lookahead, exactly as the spreadsheet does it, so
    # a month that starts before the data still resolves.
    def day_at(k):
        return DAY_NUM[rows[k][2]] if 0 <= k < len(rows) else None
    fast = (day in (8, 14, 15, 23)
            or day_at(i + 1) == 1 or day_at(i + 2) == 1)
    records.append({
        "date": d.isoformat(),
        "weekday": d.strftime("%a"),
        "monthCn": month_cn,
        "monthPy": ("run " if leap else "") + MONTH_PY[mnum - 1],
        "monthNum": mnum,
        "monthLabel": ("leap " if leap else "") + ORD[mnum] + " month",
        "leap": leap,
        "dayCn": day_cn.replace(" ", ""),
        "dayPy": DAY_PY[day - 1],
        "dayNum": day,
        "monthLength": mlen,
        "fast": bool(fast),
        "moon": "new" if day == 1 else ("full" if day == 15 else ""),
        "longFast": mnum in (1, 5, 9),
        "source": "DRBA 2026 calendar" if d <= drba_last else "HKO tables",
        "event": events.get(d, ""),
    })

meta = {
    "generated": datetime.date.today().isoformat(),
    "first": records[0]["date"],
    "last": records[-1]["date"],
    "count": len(records),
    "rules": {
        "sixFastDays": "lunar 8, 14, 15, 23 and the last two days of the "
                       "month (28+29 in a 29-day month)",
        "longFastMonths": "the 1st, 5th and 9th lunar months",
        "leapMonths": "counted as long fasting months, and flagged, pending "
                      "confirmation",
        "moon": "new moon on day 1, full moon on day 15, by the calendar "
                "convention rather than the astronomical instant",
    },
    "sources": [
        "Lunar dates: Hong Kong Observatory Gregorian-Lunar Calendar "
        "Conversion Tables, 2026-2029",
        "Dharma events and the 2026-08-12 to 2027-01-31 range: DRBA printed "
        "calendar for 2026",
        "Six fast days as 8, 14, 15, 23, 29, 30: Fo shuo si tianwang jing, "
        "Taisho T15n0590",
    ],
}

import os
os.makedirs(OUT, exist_ok=True)

print("[*] writing data.js")
with open("%s/data.js" % OUT, "w", encoding="utf-8") as f:
    f.write("// Generated by build_data.py from DRBA_Calendar.xlsx. "
            "Do not edit by hand.\n")
    f.write("window.DRBA_DATA = ")
    json.dump({"meta": meta, "days": records}, f, ensure_ascii=False)
    f.write(";\n")

print("[*] writing data.json")
with open("%s/data.json" % OUT, "w", encoding="utf-8") as f:
    json.dump({"meta": meta, "days": records}, f,
              ensure_ascii=False, indent=1)

print("[*] writing data.csv (utf-8-sig)")
cols = ["date", "weekday", "monthCn", "monthPy", "monthNum", "monthLabel",
        "leap", "dayCn", "dayPy", "dayNum", "monthLength", "fast", "moon",
        "longFast", "source", "event"]
with open("%s/data.csv" % OUT, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    for rec in records:
        w.writerow(rec)

fasts = sum(1 for r in records if r["fast"])
print("    %d days, %d fast days, %d long-fast days, %d leap days"
      % (len(records), fasts, sum(1 for r in records if r["longFast"]),
         sum(1 for r in records if r["leap"])))
print("[*] done")
